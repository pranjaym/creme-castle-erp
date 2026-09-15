import openpyxl, json, re, collections, datetime, sys
# usage: python3 workbook_extract.py <workbook.xlsx>  -> recipes_raw.json (then run workbook_normalise.py -> model.json)
SRC = sys.argv[1] if len(sys.argv) > 1 else '/Users/Pranjay/Downloads/NP - Active Items Recipe Sheet on 17th Aug-2026 Final.xlsx'
wbf = openpyxl.load_workbook(SRC, data_only=False)
wbv = openpyxl.load_workbook(SRC, data_only=True)
def num(v):
    return v if isinstance(v,(int,float)) and not isinstance(v,bool) else None
def s(v): return v.strip() if isinstance(v,str) else v

# ---------- sub mesa price list (names + unit) ----------
sm_f, sm_v = wbf['NP Sub Mesa Recipe'], wbv['NP Sub Mesa Recipe']
mesa_list = {}
for r in range(2, sm_v.max_row+1):
    n = s(sm_v.cell(r,17).value)
    if n: mesa_list[n] = {'unit_cost_sheet': num(sm_v.cell(r,19).value), 'out_unit': s(sm_v.cell(r,20).value)}

# ---------- ingredients ----------
ing_f, ing_v = wbf['Ingredients'], wbv['Ingredients']
ingredients = {}
for r in range(2, ing_v.max_row+1):
    code = s(ing_v.cell(r,1).value)
    if not code: continue
    if code in mesa_list: continue   # sub mesa rows live in the same price list
    name = s(ing_v.cell(r,2).value)
    C = num(ing_v.cell(r,3).value); D = s(ing_v.cell(r,4).value)
    E = num(ing_v.cell(r,5).value); F = s(ing_v.cell(r,6).value)
    G = num(ing_v.cell(r,7).value); H = s(ing_v.cell(r,8).value)
    ef = ing_f.cell(r,5).value; gf = ing_f.cell(r,7).value
    # pack: purchase unit -> cost unit (E = C * k  => cost units per purchase unit = 1/k)
    pack = None
    if C and E and E!=0: pack = C/E     # cost units (kg/L/each) per purchase unit
    base = (H or '').lower()
    base_unit = 'gram' if base in ('grams','gram','gm','g') else 'millilitre' if base in ('ml','millilitre') else 'piece'
    base_per_cost_unit = 1000 if base_unit in ('gram','millilitre') else 1
    # what the recipes actually pay per base unit is G
    ingredients[code] = dict(code=code, name=name, purchase_unit=D, purchase_price=C,
        cost_unit=F, price_per_cost_unit=E, pack_cost_units_per_purchase_unit=pack,
        base_unit=base_unit, base_label=H, price_per_base=G,
        e_formula=ef if isinstance(ef,str) else None, g_formula=gf if isinstance(gf,str) else None)

# ---------- generic recipe block parser ----------
LABELS = {'total','price per unit','price per unir','yield','yeild','yield price per unit','total no. of portion',
 'number of portions','per portion cost','weight per brownie','number of pcs','price pe pc','pricing per unit',
 'selling price','packaging charge','food cost with packaging','food cost without packaging','price per pc'}
def parse_sheet(ws, name_col, status_col, code_col, qty_col, price_col, amt_col, label_col, verified_col, kind):
    recipes = collections.OrderedDict()
    order = []; dupmode = False
    for r in range(2, ws.max_row+1):
        name = s(ws.cell(r,name_col).value)
        if not name or name=='N/A': continue
        code = s(ws.cell(r,code_col).value)
        label = s(ws.cell(r,label_col).value)
        if not order or order[-1] != name:
            # a new contiguous block begins
            order.append(name)
            dupmode = name in recipes
            if dupmode: recipes[name].setdefault('duplicate_blocks', []).append(dict(rows=[], lines=[]))
        if dupmode:
            blk = recipes[name]['duplicate_blocks'][-1]; blk['rows'].append(r)
            if code: blk['lines'].append(dict(ref=code, qty=num(ws.cell(r,qty_col).value), unit=s(ws.cell(r,qty_col+1).value)))
            continue
        rec = recipes.get(name)
        if rec is None:
            rec = dict(name=name, kind=kind, status=s(ws.cell(r,status_col).value), lines=[], packaging=[], sheet={}, rows=[], verified=False, notes=[])
            recipes[name]=rec
        rec['rows'].append(r)
        if verified_col and s(ws.cell(r,verified_col).value)=='Verified': rec['verified']=True
        if code:
            q = num(ws.cell(r,qty_col).value); p = num(ws.cell(r,price_col).value); a = num(ws.cell(r,amt_col).value)
            unit = s(ws.cell(r,qty_col+1).value)
            line = dict(ref=code, qty=q, unit=unit, price_sheet=p, amount_sheet=a, row=r)
            if 'unit_cost' in rec['sheet'] or 'pricing_per_unit' in rec['sheet']:
                rec['packaging'].append(line)
            else:
                rec['lines'].append(line)
        elif label:
            l = label.lower()
            H = ws.cell(r,qty_col).value; J = num(ws.cell(r,amt_col).value); F = num(ws.cell(r,price_col).value)
            if l=='total': rec['sheet']['input_qty']=num(H); rec['sheet']['batch_cost']=J
            elif l in ('price per unit','price per unir'): rec['sheet']['ppu']=J
            elif l in ('yield','yeild'): rec['sheet']['yield']=num(H)
            elif l=='yield price per unit':
                rec['sheet']['unit_cost']=J
                if isinstance(H,str): rec['sheet']['sold_size']=H
            elif l in ('total no. of portion','number of portions','number of pcs'): rec['sheet']['portions']=num(H)
            elif l in ('per portion cost','price pe pc','price per pc'): rec['sheet']['per_piece']=J
            elif l=='pricing per unit':
                rec['sheet']['pricing_per_unit']=J
                k = num(ws.cell(r,amt_col+1).value); m = num(ws.cell(r,amt_col+3).value)
                if k is not None: rec['sheet']['cost_rounded']=k
                if m is not None: rec['sheet']['price_alt']=m
            elif l=='weight per brownie': rec['sheet']['weight_per_piece']=num(H)
            elif l=='selling price': rec['sheet']['selling_price']=F
            elif l=='packaging charge': rec['sheet']['packaging_charge']=F
            elif l=='food cost with packaging': rec['sheet']['fc_with_pack']=F
            elif l=='food cost without packaging': rec['sheet']['fc_without_pack']=F
            else: rec['notes'].append(label)
    return recipes

sub = parse_sheet(sm_v, 3, 2, 4, 8, 6, 10, 5, None, 'intermediate')
for n,rec in sub.items():
    pl = mesa_list.get(n, {})
    rec['out_unit_sheet'] = pl.get('out_unit'); rec['unit_cost_pricelist'] = pl.get('unit_cost_sheet')
cakes = parse_sheet(wbv['NP Active Cake Recipe'], 1, 3, 4, 8, 6, 10, 5, 11, 'cake')
pastry = parse_sheet(wbv['NP Active Pastry Recipe'], 1, 3, 4, 8, 6, 10, 5, 14, 'pastry')

# summary tables on the right of the FG sheets (Recipe Date etc.)
def summary(ws, first_col):
    out={}
    hdr = {s(ws.cell(1,c).value):c for c in range(first_col, ws.max_column+1) if ws.cell(1,c).value}
    for r in range(2, ws.max_row+1):
        nm = s(ws.cell(r, hdr['Item Name']).value)
        if not nm: continue
        d = ws.cell(r, hdr['Recipe Date']).value
        out[nm]=dict(recipe_date=d.date().isoformat() if isinstance(d,datetime.datetime) else d,
            cost=num(ws.cell(r,hdr['Cost']).value), box=num(ws.cell(r,hdr['Box']).value), base=num(ws.cell(r,hdr['Base']).value),
            sleeve=num(ws.cell(r,hdr['Sleeve']).value), carry_bag=num(ws.cell(r,hdr['Carry Bag']).value),
            price=num(ws.cell(r,hdr['Price']).value), pc=num(ws.cell(r,hdr['PC']).value))
    return out
cake_sum = summary(wbv['NP Active Cake Recipe'], 33)
pastry_sum = summary(wbv['NP Active Pastry Recipe'], 52)
for n,rec in cakes.items(): rec['summary']=cake_sum.get(n)
for n,rec in pastry.items(): rec['summary']=pastry_sum.get(n)

data = dict(ingredients=ingredients, intermediates=sub, finished={**cakes, **pastry}, source='NP - Active Items Recipe Sheet on 17th Aug-2026 Final.xlsx')
json.dump(data, open('recipes_raw.json','w'), default=str, indent=0)
print('ingredients', len(ingredients), 'intermediates', len(sub), 'cakes', len(cakes), 'pastries', len(pastry))
print('cake statuses', collections.Counter(r['status'] for r in cakes.values()))
print('pastry statuses', collections.Counter(r['status'] for r in pastry.values()))
print('sub statuses', collections.Counter(r['status'] for r in sub.values()))
print('FG with portions', sum(1 for r in list(cakes.values())+list(pastry.values()) if r['sheet'].get('portions')))
print('notes labels', collections.Counter(x for r in list(sub.values())+list(cakes.values())+list(pastry.values()) for x in r['notes']).most_common(10))
missing_sum = [n for n,r in {**cakes,**pastry}.items() if not r.get('summary')]
print('FG without summary row', len(missing_sum), missing_sum[:8])
# refs not resolvable
allrefs = collections.Counter()
for r in list(sub.values())+list(cakes.values())+list(pastry.values()):
    for l in r['lines']+r['packaging']:
        if l['ref'] not in ingredients and l['ref'] not in sub: allrefs[l['ref']]+=1
print('unresolved refs', allrefs.most_common(15))
