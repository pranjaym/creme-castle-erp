#!/usr/bin/env python3
"""Load the team's discount sheet ("Zomato And Swiggy Discounts.xlsx") into the
coupon module's upload register (coupons.upload + coupons.upload_cell, migration
232). One upload row per platform per sheet version; the previous upload for that
platform is superseded, never deleted.

The sheet layout (Pranjay, 19 Sep 2026): row 1 = a date label in A1 and the
column-group headings (New User Base Codes, New User Flat Off, ...), row 2 = the
per-column slot (LA/MM, UM, New User, Repeat User, ...), then one row per outlet
with the platform's restaurant id in column A, the outlet code in column B and
the construct in every other cell ("60% upto 120", "Flat 125 MOV 549").

Usage:
  python3 import_sheet.py --file "<path>.xlsx" [--zomato-sheet "Zomato Discount"] [--swiggy-sheet "Swiggy Discounts"] [--by "name"] [--dry-run]
"""
import argparse, os, re, sys
import openpyxl, psycopg2

HERE = os.path.dirname(os.path.abspath(__file__))

def env_url():
    url = os.environ.get('SPINE_DATABASE_URL')
    if url: return url
    p = os.path.join(HERE, '..', '..', '.env.local')
    for line in open(p):
        m = re.match(r'^SPINE_DATABASE_URL=(.*)$', line.strip())
        if m: return m.group(1).strip('"').strip("'")
    sys.exit('SPINE_DATABASE_URL not set')

def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    h1, h2 = rows[0], rows[1]
    label = str(h1[0]) if h1[0] is not None else 'unlabelled'
    if hasattr(h1[0], 'strftime'): label = h1[0].strftime('%d %b %Y')
    # carry each group heading right until the next one
    groups, cur = [], None
    for c in h1[2:]:
        if c not in (None, ''): cur = str(c).strip()
        groups.append(cur)
    slots = [str(c).strip() if c not in (None, '') else '' for c in h2[2:]]
    cells = []
    for r in rows[2:]:
        if r[1] in (None, ''): continue
        rid = str(r[0]).strip() if r[0] is not None else None
        outlet = str(r[1]).strip()
        for i, v in enumerate(r[2:len(slots) + 2]):
            if v in (None, ''): continue
            cells.append((outlet, rid, i + 1, groups[i], slots[i], str(v).strip()))
    return label, cells

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', required=True)
    ap.add_argument('--zomato-sheet', default='Zomato Discount')
    ap.add_argument('--swiggy-sheet', default='Swiggy Discounts')
    ap.add_argument('--by', default='import_sheet.py')
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.file, data_only=True)
    plan = [('zomato', a.zomato_sheet), ('swiggy', a.swiggy_sheet)]
    cn = None if a.dry_run else psycopg2.connect(env_url(), connect_timeout=20, keepalives=1, keepalives_idle=30, keepalives_interval=10, keepalives_count=5)
    for platform, sheet in plan:
        if sheet not in wb.sheetnames:
            print(f'{platform}: sheet "{sheet}" not in workbook, skipped'); continue
        label, cells = read_sheet(wb[sheet])
        outlets = sorted(set(c[0] for c in cells))
        print(f'{platform}: "{label}", {len(outlets)} outlets, {len(cells)} cells')
        if a.dry_run:
            for c in cells[:6]: print('   ', c)
            continue
        cur = cn.cursor()
        # outlet codes that the outlet master does not know are kept as written and reported
        cur.execute('select internal_code from public.outlets')
        known = {r[0] for r in cur.fetchall()}
        unknown = [o for o in outlets if o not in known]
        if unknown: print(f'   not in public.outlets (kept as written): {unknown}')
        cur.execute("update coupons.upload set superseded_at = now() where platform = %s and superseded_at is null", (platform,))
        cur.execute("insert into coupons.upload (platform, label, source, uploaded_by, row_count) values (%s,%s,%s,%s,%s) returning id",
                    (platform, label, os.path.basename(a.file), a.by, len(outlets)))
        uid = cur.fetchone()[0]
        cur.executemany("insert into coupons.upload_cell (upload_id, outlet_code, platform_rid, col_no, slot_group, slot, construct, construct_norm) values (%s,%s,%s,%s,%s,%s,%s, coupons.norm_construct(%s))",
                        [(uid, o, rid, col, g, s, v, v) for (o, rid, col, g, s, v) in cells])
        cur.execute("insert into coupons.event (entity, entity_id, action, actor, data) values ('upload', %s, 'import', %s, %s::jsonb)",
                    (uid, a.by, '{"platform":"%s","label":"%s","cells":%d}' % (platform, label.replace('"', ''), len(cells))))
        cn.commit(); print(f'   upload id {uid} committed')
    if cn: cn.close()

if __name__ == '__main__':
    main()
